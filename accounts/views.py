from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from django.contrib.auth import get_user_model
from django.urls import reverse
import json
from django.db import models

from notices.models import Notice
from courses.models import Course
from results.models import Result
from payments.models import Payment
from accounts.models import CustomUser
from accounts.forms import (
    AdminStudentForm,
    AdminTeacherForm,
    CourseForm
)

User = get_user_model()


# ==================== LOGOUT ====================
def logout_view(request):
    logout(request)
    return redirect('login')


# ==================== GRADE HELPER ====================
def _calculate_grade_gpa(total_score):
    total = float(total_score)
    if total >= 80:
        return "A+", 4.00
    elif total >= 75:
        return "A", 3.75
    elif total >= 70:
        return "A-", 3.50
    elif total >= 65:
        return "B+", 3.25
    elif total >= 60:
        return "B", 3.00
    elif total >= 55:
        return "B-", 2.75
    elif total >= 50:
        return "C+", 2.50
    elif total >= 45:
        return "C", 2.25
    elif total >= 40:
        return "D", 2.00
    else:
        return "F", 0.00


# ==================== LOGIN & HOME ====================
def login_view(request):
    role = request.GET.get('role', 'student')
    if request.method == "POST":
        login_input = request.POST.get("login_input", "").strip()
        password = request.POST.get("password", "")

        if not login_input or not password:
            messages.error(request, "Please provide both username/email and password.")
            return render(request, "login/login.html", {'role': role})

        user = authenticate(request, username=login_input, password=password)
        if user is None:
            try:
                user_obj = User.objects.get(email=login_input)
                user = authenticate(request, username=user_obj.username, password=password)
            except User.DoesNotExist:
                pass

        if user and user.is_active:
            login(request, user)
            if user.role == "admin":
                return redirect("admin_dashboard")
            elif user.role == "teacher":
                return redirect("teacher_dashboard")
            else:
                return redirect("student_dashboard")

        messages.error(request, "Invalid username/email or password")

    return render(request, "login/login.html", {'role': role})


def home(request):
    return render(request, 'home.html')


# ==================== TEACHER VIEWS ====================
@login_required(login_url='/login/')
def teacher_dashboard(request):
    courses = Course.objects.all().order_by('course_code')
    return render(request, "teacher/dashboard.html", {"courses": courses})


@login_required(login_url='/login/')
def teacher_courses(request):
    courses = Course.objects.all().order_by('course_code')
    return render(request, 'teacher/courses.html', {'courses': courses})


@login_required(login_url='/login/')
def teacher_students(request):
    students = User.objects.filter(role='student')
    return render(request, "teacher/students_list.html", {"students": students})


@login_required(login_url='/login/')
def student_details(request, student_id):
    student = get_object_or_404(User, id=student_id, role='student')
    from students.models import StudentProfile

    profile = StudentProfile.objects.filter(user=student).first()
    enrolled_courses = Result.objects.filter(student=student).select_related('course').order_by(
        'semester', 'course__course_code'
    )
    results = Result.objects.filter(student=student).select_related('course').order_by(
        '-semester', 'course__course_code'
    )

    return render(request, "teacher/student_details.html", {
        "student": student,
        "profile": profile,
        "enrolled_courses": enrolled_courses,
        "results": results,
    })


@login_required(login_url='/login/')
def teacher_upload_marks(request):
    teacher_courses = Course.objects.filter(teacher=request.user)
    selected_course_id = request.GET.get('course')
    selected_semester = request.GET.get('semester', 1)

    try:
        selected_semester_int = int(selected_semester)
    except (ValueError, TypeError):
        selected_semester_int = 1

    if not selected_course_id and teacher_courses.exists():
        selected_course_id = str(teacher_courses.first().id)

    selected_course = None
    if selected_course_id:
        try:
            selected_course = teacher_courses.get(id=int(selected_course_id))
        except:
            selected_course = None

    students = CustomUser.objects.filter(role='student')

    if selected_course:
        existing_results = {
            r.student_id: r for r in Result.objects.filter(
                course=selected_course, semester=selected_semester_int
            )
        }
        for student in students:
            res = existing_results.get(student.id)
            if res:
                student.existing_quiz = res.quiz_marks
                student.existing_forum = res.forum_marks
                student.existing_assignment = res.assignment_marks
                student.existing_midterm = res.midterm_marks
                student.existing_final = res.final_marks
                student.existing_total = res.marks
            else:
                student.existing_quiz = 0.0
                student.existing_forum = 0.0
                student.existing_assignment = 0.0
                student.existing_midterm = 0.0
                student.existing_final = 0.0
                student.existing_total = 0.0

    if request.method == "POST":
        course_id = request.POST.get("course")
        semester = int(request.POST.get("semester", 1))
        course = get_object_or_404(Course, id=course_id, teacher=request.user)

        for student in students:
            quiz = float(request.POST.get(f"quiz_{student.id}", 0) or 0)
            forum = float(request.POST.get(f"forum_{student.id}", 0) or 0)
            assignment = float(request.POST.get(f"assignment_{student.id}", 0) or 0)
            midterm = float(request.POST.get(f"midterm_{student.id}", 0) or 0)
            final = float(request.POST.get(f"final_{student.id}", 0) or 0)

            total = quiz + forum + assignment + midterm + final
            if total == 0:
                continue

            grade, gpa = _calculate_grade_gpa(total)

            Result.objects.update_or_create(
                student=student,
                course=course,
                semester=semester,
                defaults={
                    "quiz_marks": quiz,
                    "forum_marks": forum,
                    "assignment_marks": assignment,
                    "midterm_marks": midterm,
                    "final_marks": final,
                    "marks": total,
                    "grade": grade,
                    "gpa": gpa,
                }
            )

        messages.success(request, "Marks uploaded successfully.")
        return redirect(reverse('teacher_upload_marks') + f'?course={course_id}&semester={semester}')

    return render(request, "teacher/upload_marks.html", {
        "students": students,
        "courses": teacher_courses,
        "semesters": range(1, 13),
        "selected_course": selected_course,
        "selected_semester": selected_semester_int,
    })


@login_required(login_url='/login/')
def teacher_results(request):
    selected_semester = request.GET.get('semester')
    teacher_courses = Course.objects.filter(teacher=request.user)
    results = Result.objects.filter(course__in=teacher_courses)

    if selected_semester:
        try:
            results = results.filter(semester=int(selected_semester))
        except:
            pass

    results = results.select_related('student', 'course').order_by('course__course_code', 'student__username')

    return render(request, 'teacher/results.html', {
        'results': results,
        'semesters': range(1, 13),
        'selected_semester': selected_semester or "",
    })


# ==================== STUDENT VIEWS ====================
@login_required(login_url='/login/')
def student_dashboard(request):
    if getattr(request.user, 'role', None) != 'student':
        return redirect('home')

    results = Result.objects.filter(student=request.user).select_related('course').order_by('semester')

    semester_map = {}
    for r in results:
        gpa_value = float(r.gpa or 0)
        semester_map.setdefault(r.semester, []).append(gpa_value)

    labels = [f"Sem {sem}" for sem in sorted(semester_map.keys())]
    gpas = [round(sum(vals) / len(vals), 2) if vals else 0.00 for vals in semester_map.values()]

    payments = Payment.objects.filter(student=request.user).order_by('-date')

    result_semesters = sorted(set(results.values_list('semester', flat=True)))
    payment_semesters = sorted(set(payments.values_list('semester', flat=True)))
    all_semesters = sorted(set(result_semesters + payment_semesters))
    notices = Notice.objects.filter(is_active=True).order_by('-created_at')[:5]

    return render(request, "student/dashboard.html", {
        "semesters_json": json.dumps(labels),
        "gpas_json": json.dumps(gpas),
        "payments": payments,
        "semesters": all_semesters,
        "has_results": results.exists(),
        "has_payments": payments.exists(),
        "results_count": results.count(),
        "student_id": request.user.id,
        "notices": notices,
    })


@login_required(login_url='/login/')
def student_results(request):
    if request.user.role != 'student':
        return redirect('home')

    results = Result.objects.filter(student=request.user).select_related('course').order_by('semester')
    total_gpa = sum(float(r.gpa or 0) for r in results)
    cgpa = round(total_gpa / results.count(), 2) if results.exists() else 0.00

    semesters = sorted(list(set(results.values_list('semester', flat=True))))
    current_semester = max(semesters) if semesters else None
    selected_semester = request.GET.get('semester', 'current')

    if selected_semester == 'all':
        current_results = results
        previous_results = {}
    elif selected_semester == 'current':
        current_results = results.filter(semester=current_semester) if current_semester else results.none()
        previous_results = {}
        for sem in semesters:
            if sem != current_semester:
                previous_results[f"Semester {sem}"] = results.filter(semester=sem)
    else:
        try:
            semester_no = int(selected_semester)
            current_results = results.filter(semester=semester_no)
            previous_results = {}
        except:
            current_results = results.filter(semester=current_semester) if current_semester else results.none()
            previous_results = {}

    return render(request, "student/results.html", {
        "cgpa": cgpa,
        "current_semester": current_semester,
        "current_results": current_results,
        "previous_results": previous_results,
        "semesters": semesters,
        "selected_semester": selected_semester,
    })


@login_required(login_url='/login/')
def student_courses(request):
    if request.user.role != 'student':
        return redirect('home')

    results = Result.objects.filter(student=request.user).select_related('course').order_by('semester', 'course__course_code')
    semesters = sorted(list(set(results.values_list('semester', flat=True))))
    current_semester = max(semesters) if semesters else None
    selected_semester = request.GET.get('semester', 'current')

    if selected_semester == 'all':
        filtered_results = results
    elif selected_semester == 'current':
        filtered_results = results.filter(semester=current_semester)
    else:
        try:
            filtered_results = results.filter(semester=int(selected_semester))
        except:
            filtered_results = results.filter(semester=current_semester)

    grouped_courses = {}
    for result in filtered_results:
        grouped_courses.setdefault(result.semester, []).append(result)

    return render(request, "student/courses.html", {
        "grouped_courses": grouped_courses,
        "semesters": semesters,
        "current_semester": current_semester,
        "selected_semester": selected_semester,
    })


@login_required(login_url='/login/')
def student_profile(request):
    if request.user.role != 'student':
        return redirect('home')

    if request.method == 'POST':
        request.user.phone = request.POST.get('phone')
        request.user.address = request.POST.get('address')
        if request.FILES.get('profile_picture'):
            request.user.profile_picture = request.FILES['profile_picture']
        request.user.save()
        messages.success(request, 'Profile updated successfully.')
        return redirect('student_profile')

    return render(request, "student/profile.html", {"user": request.user})


# ==================== ADMIN VIEWS ====================
@login_required(login_url='/login/')
def admin_dashboard(request):
    if request.user.role != 'admin':
        return redirect('home')
    
    total_students = User.objects.filter(role='student').count()
    total_teachers = User.objects.filter(role='teacher').count()
    total_courses = Course.objects.count()

    students = User.objects.filter(role='student').order_by('username')[:5]
    courses = Course.objects.all().order_by('course_code')[:5]

    return render(request, 'admin/dashboard.html', {
        'total_students': total_students,
        'total_teachers': total_teachers,
        'total_courses': total_courses,
        'students': students,
        'courses': courses,
    })


@login_required(login_url='/login/')
def admin_students(request):
    if request.user.role != 'admin':
        return redirect('home')
    
    search_query = request.GET.get('search', '').strip()
    selected_semester = request.GET.get('semester', '')
    
    students = User.objects.filter(role='student')
    
    if search_query:
        students = students.filter(
            models.Q(username__icontains=search_query) |
            models.Q(first_name__icontains=search_query) |
            models.Q(last_name__icontains=search_query) |
            models.Q(email__icontains=search_query) |
            models.Q(id__icontains=search_query)
        )
        
    if selected_semester:
        try:
            sem_val = int(selected_semester)
            students = students.filter(
                models.Q(results__semester=sem_val) | models.Q(payments__semester=sem_val)
            ).distinct()
        except ValueError:
            pass
            
    students = students.order_by('username')
    
    payment_semesters = list(Payment.objects.values_list('semester', flat=True).distinct())
    result_semesters = list(Result.objects.values_list('semester', flat=True).distinct())
    semesters = sorted(list(set(payment_semesters + result_semesters)))
    if not semesters:
        semesters = list(range(1, 13))
        
    return render(request, 'admin/students.html', {
        'students': students,
        'semesters': semesters,
        'selected_semester': int(selected_semester) if selected_semester.isdigit() else '',
        'search_query': search_query,
    })


@login_required(login_url='/login/')
def admin_teachers(request):
    if request.user.role != 'admin':
        return redirect('home')
    teachers = User.objects.filter(role='teacher').order_by('username')
    return render(request, 'admin/teachers.html', {'teachers': teachers})


@login_required(login_url='/login/')
def admin_courses(request):
    if request.user.role != 'admin':
        return redirect('home')
        
    search_query = request.GET.get('search', '').strip()
    courses = Course.objects.all()
    
    if search_query:
        courses = courses.filter(
            models.Q(course_code__icontains=search_query) |
            models.Q(course_title__icontains=search_query)
        )
        
    courses = courses.order_by('course_code')
    return render(request, 'admin/courses.html', {
        'courses': courses,
        'search_query': search_query,
    })


@login_required(login_url='/login/')
def admin_results(request):
    if request.user.role != 'admin':
        return redirect('home')
    
    search_query = request.GET.get('search', '').strip()
    selected_semester = request.GET.get('semester')

    results = Result.objects.select_related('student', 'course')

    if search_query:
        results = results.filter(
            models.Q(student__username__icontains=search_query) |
            models.Q(student__first_name__icontains=search_query) |
            models.Q(student__last_name__icontains=search_query) |
            models.Q(course__course_code__icontains=search_query)
        )

    if selected_semester:
        try:
            results = results.filter(semester=int(selected_semester))
        except (ValueError, TypeError):
            pass

    results = results.order_by('-semester', 'student__username')

    semesters = sorted(set(Result.objects.values_list('semester', flat=True).distinct()))

    return render(request, 'admin/results.html', {
        'results': results,
        'semesters': semesters,
        'selected_semester': selected_semester,
        'search_query': search_query,
    })


@login_required(login_url='/login/')
def admin_payments(request):
    if request.user.role != 'admin':
        return redirect('home')
    
    search_query = request.GET.get('search', '').strip()
    selected_semester = request.GET.get('semester', '')
    
    try:
        payments = Payment.objects.select_related('student')
        
        if search_query:
            payments = payments.filter(
                models.Q(student__username__icontains=search_query) |
                models.Q(student__first_name__icontains=search_query) |
                models.Q(student__last_name__icontains=search_query) |
                models.Q(amount__icontains=search_query)
            )
            
        if selected_semester:
            try:
                payments = payments.filter(semester=int(selected_semester))
            except ValueError:
                pass
                
        payments = payments.order_by('-date')
        semesters = sorted(list(set(Payment.objects.values_list('semester', flat=True).distinct())))
    except:
        payments = []
        semesters = list(range(1, 13))

    return render(request, 'admin/payments.html', {
        'payments': payments,
        'search_query': search_query,
        'semesters': semesters,
        'selected_semester': int(selected_semester) if selected_semester.isdigit() else '',
    })


@login_required(login_url='/login/')
def admin_payment_add(request, student_id):
    if request.user.role != 'admin':
        return redirect('home')

    student = get_object_or_404(CustomUser, id=student_id, role='student')

    if request.method == 'POST':
        Payment.objects.create(
            student=student,
            semester=int(request.POST.get('semester', 1)),
            amount=request.POST.get('amount', 0),
            status=request.POST.get('status', 'Pending'),
            invoice_id=request.POST.get('invoice_id', ''),
        )
        messages.success(request, 'Payment record added successfully.')
        return redirect('admin_student_details', student_id=student_id)

    return render(request, 'admin/payment_form.html', {
        'student': student,
        'payment': None,
        'semesters': range(1, 13),
    })


@login_required(login_url='/login/')
def admin_payment_edit(request, payment_id):
    if request.user.role != 'admin':
        return redirect('home')

    payment = get_object_or_404(Payment, id=payment_id)
    student = payment.student

    if request.method == 'POST':
        payment.semester = int(request.POST.get('semester', payment.semester))
        payment.amount = request.POST.get('amount', payment.amount)
        payment.status = request.POST.get('status', payment.status)
        payment.invoice_id = request.POST.get('invoice_id', payment.invoice_id or '')
        payment.save()
        messages.success(request, 'Payment record updated successfully.')
        return redirect('admin_payments')

    return render(request, 'admin/payment_form.html', {
        'student': student,
        'payment': payment,
        'semesters': range(1, 13),
    })


@login_required(login_url='/login/')
def admin_payment_delete(request, payment_id):
    if request.user.role != 'admin':
        return redirect('home')

    payment = get_object_or_404(Payment, id=payment_id)
    student_id = payment.student_id
    payment.delete()
    messages.success(request, 'Payment record deleted successfully.')
    return redirect('admin_payments')


@login_required(login_url='/login/')
def admin_payment_add_direct(request):
    if request.user.role != 'admin':
        return redirect('home')

    students = CustomUser.objects.filter(role='student').order_by('username')

    if request.method == 'POST':
        student_id = request.POST.get('student')
        student = get_object_or_404(CustomUser, id=student_id, role='student')
        
        Payment.objects.create(
            student=student,
            semester=int(request.POST.get('semester', 1)),
            amount=request.POST.get('amount', 0),
            status=request.POST.get('status', 'Pending'),
            invoice_id=request.POST.get('invoice_id', ''),
        )
        messages.success(request, 'Payment record added successfully.')
        return redirect('admin_payments')

    return render(request, 'admin/payment_form_direct.html', {
        'students': students,
        'semesters': range(1, 13),
    })


@login_required(login_url='/login/')
def admin_upload_marks(request):
    if request.user.role != 'admin':
        return redirect('home')

    students = CustomUser.objects.filter(role='student')
    courses = Course.objects.all()
    semesters = range(1, 9)

    selected_course_id = request.GET.get('course')
    selected_semester = request.GET.get('semester', 1)

    if not selected_course_id and courses.exists():
        selected_course_id = str(courses.first().id)

    selected_course = None
    if selected_course_id:
        try:
            selected_course = courses.get(id=int(selected_course_id))
        except:
            selected_course = None

    if selected_course:
        try:
            selected_semester_int = int(selected_semester)
        except (ValueError, TypeError):
            selected_semester_int = 1
            
        existing_results = {
            r.student_id: r for r in Result.objects.filter(
                course=selected_course, semester=selected_semester_int
            )
        }
        for student in students:
            res = existing_results.get(student.id)
            if res:
                student.existing_quiz = res.quiz_marks
                student.existing_forum = res.forum_marks
                student.existing_assignment = res.assignment_marks
                student.existing_midterm = res.midterm_marks
                student.existing_final = res.final_marks
                student.existing_total = res.marks
            else:
                student.existing_quiz = 0.0
                student.existing_forum = 0.0
                student.existing_assignment = 0.0
                student.existing_midterm = 0.0
                student.existing_final = 0.0
                student.existing_total = 0.0

    if request.method == "POST":
        course_id = request.POST.get("course")
        semester = int(request.POST.get("semester", 1))
        course = get_object_or_404(Course, id=course_id)

        for student in students:
            quiz = float(request.POST.get(f"quiz_{student.id}", 0) or 0)
            forum = float(request.POST.get(f"forum_{student.id}", 0) or 0)
            assignment = float(request.POST.get(f"assignment_{student.id}", 0) or 0)
            midterm = float(request.POST.get(f"midterm_{student.id}", 0) or 0)
            final = float(request.POST.get(f"final_{student.id}", 0) or 0)

            total = quiz + forum + assignment + midterm + final
            if total == 0:
                continue

            grade, gpa = _calculate_grade_gpa(total)

            Result.objects.update_or_create(
                student=student,
                course=course,
                semester=semester,
                defaults={
                    "quiz_marks": quiz,
                    "forum_marks": forum,
                    "assignment_marks": assignment,
                    "midterm_marks": midterm,
                    "final_marks": final,
                    "marks": total,
                    "grade": grade,
                    "gpa": gpa,
                }
            )

        messages.success(request, "Marks uploaded successfully.")
        return redirect(reverse('admin_upload_marks') + f'?course={course_id}&semester={semester}')

    return render(request, "admin/upload_marks.html", {
        "students": students,
        "courses": courses,
        "semesters": semesters,
        "selected_course": selected_course,
        "selected_semester": int(selected_semester),
    })


# ==================== ALL REMAINING VIEWS ====================
# (Student CRUD, Teacher CRUD, Course CRUD, Notices, Excel, etc.)

@login_required(login_url='/login/')
def admin_student_add(request):
    if request.method == "POST":
        form = AdminStudentForm(request.POST)
        if form.is_valid():
            user = form.save(commit=False)
            password = form.cleaned_data.get("password")
            if password:
                user.set_password(password)
            user.role = "student"
            user.save()
            messages.success(request, "Student added successfully.")
            return redirect("admin_students")
    else:
        form = AdminStudentForm(initial={"role": "student"})
    return render(request, "admin/student_add.html", {"form": form})


@login_required(login_url='/login/')
def admin_student_edit(request, id):
    student = get_object_or_404(User, id=id, role="student")
    if request.method == "POST":
        form = AdminStudentForm(request.POST, instance=student)
        if form.is_valid():
            user = form.save(commit=False)
            password = form.cleaned_data.get("password")
            if password:
                user.set_password(password)
            user.save()
            messages.success(request, "Student updated successfully.")
            return redirect("admin_students")
    else:
        form = AdminStudentForm(instance=student)
    return render(request, "admin/student_edit.html", {"form": form, "student": student})


@login_required(login_url='/login/')
def admin_student_delete(request, id):
    student = get_object_or_404(User, id=id, role="student")
    if request.method == "POST":
        student.delete()
        messages.success(request, "Student deleted successfully.")
        return redirect("admin_students")
    return render(request, "admin/confirm_delete.html", {
        "object": student,
        "cancel_url": reverse("admin_students"),
    })


@login_required(login_url='/login/')
def admin_teacher_add(request):
    if request.method == "POST":
        form = AdminTeacherForm(request.POST)
        if form.is_valid():
            user = form.save(commit=False)
            password = form.cleaned_data.get("password")
            if password:
                user.set_password(password)
            user.role = "teacher"
            user.save()
            messages.success(request, "Teacher added successfully.")
            return redirect("admin_teachers")
    else:
        form = AdminTeacherForm(initial={"role": "teacher"})
    return render(request, "admin/teacher_add.html", {"form": form})


@login_required(login_url='/login/')
def admin_teacher_edit(request, id):
    teacher = get_object_or_404(User, id=id, role="teacher")
    if request.method == "POST":
        form = AdminTeacherForm(request.POST, instance=teacher)
        if form.is_valid():
            user = form.save(commit=False)
            password = form.cleaned_data.get("password")
            if password:
                user.set_password(password)
            user.save()
            messages.success(request, "Teacher updated successfully.")
            return redirect("admin_teachers")
    else:
        form = AdminTeacherForm(instance=teacher)
    return render(request, "admin/teacher_edit.html", {"form": form, "teacher": teacher})


@login_required(login_url='/login/')
def admin_teacher_delete(request, id):
    teacher = get_object_or_404(User, id=id, role="teacher")
    if request.method == "POST":
        teacher.delete()
        messages.success(request, "Teacher deleted successfully.")
        return redirect("admin_teachers")
    return render(request, "admin/confirm_delete.html", {
        'object': teacher,
        'cancel_url': reverse('admin_teachers'),
        'title': 'Delete Teacher'
    })


@login_required(login_url='/login/')
def admin_teacher_details(request, id):
    teacher = get_object_or_404(CustomUser, id=id, role='teacher')
    courses = Course.objects.filter(teacher=teacher)
    return render(request, "admin/teacher_details.html", {
        "teacher": teacher,
        "courses": courses
    })


@login_required(login_url='/login/')
def admin_course_add(request):
    if request.method == "POST":
        form = CourseForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, "Course added successfully.")
            return redirect("admin_courses")
    else:
        form = CourseForm()
    return render(request, "admin/course_add.html", {"form": form})


@login_required(login_url='/login/')
def admin_course_edit(request, id):
    course = get_object_or_404(Course, id=id)
    if request.method == "POST":
        form = CourseForm(request.POST, instance=course)
        if form.is_valid():
            form.save()
            messages.success(request, "Course updated successfully.")
            return redirect("admin_courses")
    else:
        form = CourseForm(instance=course)
    return render(request, "admin/course_edit.html", {"form": form, "course": course})


@login_required(login_url='/login/')
def admin_course_delete(request, id):
    course = get_object_or_404(Course, id=id)
    if request.method == "POST":
        course.delete()
        messages.success(request, "Course deleted successfully.")
        return redirect("admin_courses")
    return render(request, "admin/confirm_delete.html", {
        "object": course,
        "cancel_url": reverse("admin_courses"),
    })


@login_required(login_url='/login/')
def assign_course(request):
    if request.user.role != 'admin':
        return redirect('home')

    teachers = User.objects.filter(role='teacher')
    courses = Course.objects.all()

    if request.method == "POST":
        teacher_id = request.POST.get("teacher")
        course_id = request.POST.get("course")
        teacher = get_object_or_404(User, id=teacher_id, role='teacher')
        course = get_object_or_404(Course, id=course_id)
        course.teacher = teacher
        course.save()
        messages.success(request, "Course assigned successfully!")
        return redirect('admin_courses')

    return render(request, "admin/assign_course.html", {
        "teachers": teachers,
        "courses": courses
    })


@login_required(login_url='/login/')
def assign_course_student(request):
    if request.user.role != 'admin':
        return redirect('home')

    students = User.objects.filter(role='student')
    courses = Course.objects.all()

    if request.method == "POST":
        student_id = request.POST.get("student")
        course_id = request.POST.get("course")
        semester = request.POST.get("semester", 1)

        student = get_object_or_404(User, id=student_id, role='student')
        course = get_object_or_404(Course, id=course_id)

        try:
            semester_int = int(semester)
        except (ValueError, TypeError):
            semester_int = 1

        _, created = Result.objects.get_or_create(
            student=student,
            course=course,
            semester=semester_int,
        )

        if created:
            messages.success(request, "Course assigned to student successfully!")
        else:
            messages.warning(
                request,
                "Student is already enrolled in this course for the selected semester."
            )
        return redirect('admin_courses')

    return render(request, "admin/assign_course_student.html", {
        "students": students,
        "courses": courses,
    })


@login_required(login_url='/login/')
def admin_student_details(request, student_id):
    student = get_object_or_404(CustomUser, id=student_id)
    all_results = Result.objects.filter(student=student).order_by("-semester")
    semesters = Result.objects.filter(student=student).values_list("semester", flat=True).distinct().order_by('semester')
    current_semester = semesters.last() if semesters else None
    semester_param = request.GET.get("semester", "current")

    if semester_param == "all":
        results = all_results
    elif semester_param == "current":
        results = all_results.filter(semester=current_semester) if current_semester else all_results.none()
    else:
        try:
            results = all_results.filter(semester=int(semester_param))
        except:
            results = all_results.filter(semester=current_semester) if current_semester else all_results.none()

    payments = Payment.objects.filter(student=student).order_by('-date')
    payment_semesters = sorted(set(payments.values_list('semester', flat=True)))

    return render(request, "admin/student_details.html", {
        "student": student,
        "results": results,
        "payments": payments,
        "payment_semesters": payment_semesters,
        "semesters": semesters,
        "selected_semester": semester_param,
        "current_semester": current_semester
    })


@login_required(login_url='/login/')
def admin_notices(request):
    notices = Notice.objects.all().order_by('-created_at')
    return render(request, 'admin/notices.html', {'notices': notices})


@login_required(login_url='/login/')
def admin_notice_add(request):
    if request.method == "POST":
        Notice.objects.create(
            title=request.POST.get("title"),
            message=request.POST.get("message"),
            created_by=request.user
        )
        messages.success(request, "Notice added successfully.")
        return redirect("admin_notices")
    return render(request, "admin/notice_form.html")


@login_required(login_url='/login/')
def admin_notice_edit(request, id):
    notice = get_object_or_404(Notice, id=id)
    if request.method == "POST":
        notice.title = request.POST.get("title")
        notice.message = request.POST.get("message")
        notice.is_active = request.POST.get("is_active") == "on"
        notice.save()
        messages.success(request, "Notice updated successfully.")
        return redirect("admin_notices")
    return render(request, "admin/notice_form.html", {"notice": notice})


@login_required(login_url='/login/')
def admin_notice_delete(request, id):
    notice = get_object_or_404(Notice, id=id)
    if request.method == "POST":
        notice.delete()
        messages.success(request, "Notice deleted successfully.")
        return redirect("admin_notices")
    return render(request, "admin/confirm_delete.html", {
        "object": notice,
        "cancel_url": reverse("admin_notices"),
    })


# ==================== EXCEL UPLOADS ====================
from openpyxl import load_workbook, Workbook
from django.http import HttpResponse


@login_required(login_url='/login/')
def teacher_upload_excel(request):
    courses = Course.objects.filter(teacher=request.user)
    selected_semester = request.GET.get("semester") or request.POST.get("semester", "")

    if request.method == "POST":
        course_id = request.POST.get("course")
        semester = request.POST.get("semester")
        file = request.FILES.get("excel_file")

        if not course_id or not semester or not file:
            messages.error(request, "Course, Semester and File required.")
            return redirect("teacher_upload_excel")

        course = get_object_or_404(Course, id=course_id)
        semester = int(semester)

        wb = load_workbook(file)
        ws = wb.active

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
            try:
                student = CustomUser.objects.get(id=row[0], role='student')
            except:
                continue

            quiz = float(row[2] or 0)
            forum = float(row[3] or 0)
            assignment = float(row[4] or 0)
            midterm = float(row[5] or 0)
            final = float(row[6] or 0)
            total = quiz + forum + assignment + midterm + final

            Result.objects.update_or_create(
                student=student,
                course=course,
                semester=semester,
                defaults={
                    "quiz_marks": quiz,
                    "forum_marks": forum,
                    "assignment_marks": assignment,
                    "midterm_marks": midterm,
                    "final_marks": final,
                    "marks": total,
                }
            )

        messages.success(request, "Excel uploaded successfully!")
        return redirect("teacher_upload_excel")

    return render(request, "teacher/upload_excel.html", {
        "courses": courses,
        "semesters": range(1, 13),
        "selected_semester": selected_semester,
    })


@login_required(login_url='/login/')
def admin_upload_excel(request):
    courses = Course.objects.all()
    selected_semester = request.GET.get("semester") or request.POST.get("semester", "")

    if request.method == "POST":
        course_id = request.POST.get("course")
        semester = request.POST.get("semester")
        file = request.FILES.get("excel_file")

        if not file:
            messages.error(request, "Please select an Excel file.")
            return redirect("admin_upload_excel")
        if not semester or not course_id:
            messages.error(request, "Please select semester and course.")
            return redirect("admin_upload_excel")

        semester = int(semester)
        course = get_object_or_404(Course, id=course_id)

        try:
            wb = load_workbook(file)
            ws = wb.active
            success_count = 0
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or not row[0]:
                    continue
                try:
                    student = CustomUser.objects.get(id=row[0], role='student')
                except:
                    continue

                quiz = float(row[2] or 0)
                forum = float(row[3] or 0)
                assignment = float(row[4] or 0)
                midterm = float(row[5] or 0)
                final = float(row[6] or 0)
                total = quiz + forum + assignment + midterm + final

                Result.objects.update_or_create(
                    student=student,
                    course=course,
                    semester=semester,
                    defaults={
                        "quiz_marks": quiz,
                        "forum_marks": forum,
                        "assignment_marks": assignment,
                        "midterm_marks": midterm,
                        "final_marks": final,
                        "marks": total,
                    }
                )
                success_count += 1

            messages.success(request, f"{success_count} student marks uploaded successfully!")
            return redirect("admin_upload_excel")
        except Exception as e:
            messages.error(request, f"Error processing file: {str(e)}")
            return redirect("admin_upload_excel")

    return render(request, "admin/upload_excel.html", {
        "courses": courses,
        "selected_semester": selected_semester,
        "semesters": list(range(1, 13)),
    })


@login_required(login_url='/login/')
def download_excel_template(request):
    wb = Workbook()
    ws = wb.active
    ws.title = "Marks Template"
    ws.append(["student_id", "student_name", "quiz", "forum", "assignment", "midterm", "final"])

    students = CustomUser.objects.filter(role='student').order_by('id')
    for student in students:
        ws.append([student.id, student.username, "", "", "", "", ""])

    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = "attachment; filename=marks_template.xlsx"
    wb.save(response)
    return response